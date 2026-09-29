from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
import pytest

from concilia_fiscal import FileKind, ValidationError, validate_file

FIXTURES = Path(__file__).parent / "fixtures"


def test_validates_accounting_fixture() -> None:
    result = validate_file(
        FIXTURES / "contabil_teste.xlsx",
        file_name="contabil_teste.xlsx",
        file_kind=FileKind.ACCOUNTING,
    )

    assert result.is_valid
    assert result.row_count == 12
    assert result.recognized_columns == (
        "Código",
        "Histórico",
        "D/C",
        "Vlr Saldo Final",
        "CNPJ",
        "Data",
        "Número Lançamento",
    )
    assert result.extra_columns == ()
    assert result.errors == ()
    assert result.data is not None


@pytest.mark.parametrize(
    ("file_name", "file_kind", "expected_rows"),
    [
        ("fiscal_teste.xlsx", FileKind.FISCAL, 11),
        ("plano_contas_teste.xlsx", FileKind.ACCOUNTS, 6),
    ],
)
def test_validates_other_input_fixtures(
    file_name: str,
    file_kind: FileKind,
    expected_rows: int,
) -> None:
    result = validate_file(
        FIXTURES / file_name,
        file_name=file_name,
        file_kind=file_kind,
    )

    assert result.is_valid
    assert result.row_count == expected_rows
    assert result.errors == ()


def _save_changed_workbook(
    tmp_path: Path,
    source_name: str,
    change: Any,
) -> Path:
    workbook = openpyxl.load_workbook(FIXTURES / source_name)
    change(workbook.active)
    output = tmp_path / source_name
    workbook.save(output)
    return output


def test_accepts_and_reports_extra_columns(tmp_path: Path) -> None:
    def add_extra_column(sheet: Any) -> None:
        sheet.cell(1, 8, "Observação")
        sheet.cell(2, 8, "importada")

    source = _save_changed_workbook(tmp_path, "contabil_teste.xlsx", add_extra_column)

    result = validate_file(
        source,
        file_name=source.name,
        file_kind=FileKind.ACCOUNTING,
    )

    assert result.is_valid
    assert result.extra_columns == ("Observação",)
    assert result.data is not None
    assert "Observação" in result.data.columns


def test_reports_missing_column_and_preserves_readable_data(tmp_path: Path) -> None:
    source = _save_changed_workbook(
        tmp_path,
        "fiscal_teste.xlsx",
        lambda sheet: sheet.delete_cols(4),
    )

    result = validate_file(source, file_name=source.name, file_kind=FileKind.FISCAL)

    assert not result.is_valid
    assert result.data is not None
    assert result.errors == (
        ValidationError(
            file_name="fiscal_teste.xlsx",
            line=1,
            column="Fornecedor",
            code="missing_required_column",
            message='A coluna obrigatória "Fornecedor" não foi encontrada.',
        ),
    )


def test_collects_all_safe_row_errors_with_excel_line_numbers(tmp_path: Path) -> None:
    def add_invalid_values(sheet: Any) -> None:
        sheet["A2"] = None
        sheet["E3"] = "20/01/2026"
        sheet["C4"] = 750.123

    source = _save_changed_workbook(tmp_path, "fiscal_teste.xlsx", add_invalid_values)

    result = validate_file(source, file_name=source.name, file_kind=FileKind.FISCAL)

    assert not result.is_valid
    assert [(error.line, error.column, error.code) for error in result.errors] == [
        (2, "Alíquota PIS", "missing_required_value"),
        (3, "Data Fiscal", "invalid_value_type"),
        (4, "Vlr Documento", "invalid_value_type"),
    ]
    assert result.data is not None


def test_rejects_dates_outside_exact_year_month_day_format(tmp_path: Path) -> None:
    def add_compact_iso_date(sheet: Any) -> None:
        sheet["E2"] = "20260115"

    source = _save_changed_workbook(tmp_path, "fiscal_teste.xlsx", add_compact_iso_date)

    result = validate_file(source, file_name=source.name, file_kind=FileKind.FISCAL)

    assert [(error.line, error.column, error.code) for error in result.errors] == [
        (2, "Data Fiscal", "invalid_value_type"),
    ]


def test_rejects_scientific_notation_for_decimal_values(tmp_path: Path) -> None:
    def add_scientific_notation(sheet: Any) -> None:
        sheet["A2"] = "1e2"

    source = _save_changed_workbook(tmp_path, "fiscal_teste.xlsx", add_scientific_notation)

    result = validate_file(source, file_name=source.name, file_kind=FileKind.FISCAL)

    assert [(error.line, error.column, error.code) for error in result.errors] == [
        (2, "Alíquota PIS", "invalid_value_type"),
    ]


def test_rejects_duplicate_columns(tmp_path: Path) -> None:
    def duplicate_column(sheet: Any) -> None:
        sheet["B1"] = "Código"

    source = _save_changed_workbook(tmp_path, "contabil_teste.xlsx", duplicate_column)

    result = validate_file(source, file_name=source.name, file_kind=FileKind.ACCOUNTING)

    assert not result.is_valid
    assert result.data is None
    assert result.errors[0].code == "duplicate_column"
    assert result.errors[0].line == 1
    assert result.errors[0].column == "Código"


def test_ignores_fully_empty_rows(tmp_path: Path) -> None:
    def insert_empty_row(sheet: Any) -> None:
        sheet.insert_rows(3)

    source = _save_changed_workbook(tmp_path, "contabil_teste.xlsx", insert_empty_row)

    result = validate_file(source, file_name=source.name, file_kind=FileKind.ACCOUNTING)

    assert result.is_valid
    assert result.row_count == 12
    assert result.errors == ()


@pytest.mark.parametrize(
    ("file_name", "content", "expected_code"),
    [
        ("contabil.csv", b"not,xlsx", "invalid_file_extension"),
        ("contabil.xlsx", b"not an xlsx workbook", "unreadable_workbook"),
    ],
)
def test_rejects_files_that_cannot_be_validated(
    file_name: str,
    content: bytes,
    expected_code: str,
) -> None:
    result = validate_file(content, file_name=file_name, file_kind=FileKind.ACCOUNTING)

    assert not result.is_valid
    assert result.data is None
    assert result.errors[0].code == expected_code


def test_reports_missing_expected_sheet(tmp_path: Path) -> None:
    def rename_sheet(sheet: Any) -> None:
        sheet.title = "Outra_Aba"

    source = _save_changed_workbook(tmp_path, "plano_contas_teste.xlsx", rename_sheet)

    result = validate_file(source, file_name=source.name, file_kind=FileKind.ACCOUNTS)

    assert not result.is_valid
    assert result.data is None
    assert result.errors[0].code == "missing_expected_sheet"


def test_rejects_upload_larger_than_ten_megabytes() -> None:
    content = b"x" * (10 * 1024 * 1024 + 1)

    result = validate_file(content, file_name="contabil.xlsx", file_kind=FileKind.ACCOUNTING)

    assert not result.is_valid
    assert result.data is None
    assert result.errors[0].code == "workbook_too_large"


def test_rejects_workbook_with_too_many_archive_entries() -> None:
    content = BytesIO()
    with ZipFile(content, "w", ZIP_DEFLATED) as archive:
        for index in range(1_001):
            archive.writestr(f"entry-{index}.xml", "x")

    result = validate_file(
        content.getvalue(),
        file_name="contabil.xlsx",
        file_kind=FileKind.ACCOUNTING,
    )

    assert not result.is_valid
    assert result.data is None
    assert result.errors[0].code == "workbook_too_complex"


def test_rejects_workbook_that_expands_beyond_safe_limit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr("concilia_fiscal.excel.MAX_UNCOMPRESSED_BYTES", 5)
    content = BytesIO()
    with ZipFile(content, "w", ZIP_DEFLATED) as archive:
        archive.writestr("large.xml", "123456")

    result = validate_file(
        content.getvalue(),
        file_name="contabil.xlsx",
        file_kind=FileKind.ACCOUNTING,
    )

    assert not result.is_valid
    assert result.data is None
    assert result.errors[0].code == "workbook_too_large"
